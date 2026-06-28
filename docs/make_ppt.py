#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成组会汇报 PPT(16:9,嵌入 docs/figures/ 的图)。

纯 CPU。用法:CUDA_VISIBLE_DEVICES="" uv run python docs/make_ppt.py
输出 docs/SO-ARM101_组会汇报.pptx
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(__file__)
FIG = os.path.join(HERE, "figures")
CJK = "Noto Sans CJK SC"

NAVY = RGBColor(0x2F, 0x4B, 0x7C)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
ACC = RGBColor(0x3B, 0x6E, 0xA5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC0, 0x39, 0x2B)
ORANGE = RGBColor(0xC8, 0x84, 0x3C)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set_font(run, name=CJK, size=18, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def rect(slide, l, t, w, h, color):
    sp = slide.shapes.add_shape(1, l, t, w, h)  # rectangle
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def title_bar(slide, text, num=None):
    rect(slide, 0, 0, SW, Inches(1.0), NAVY)
    tf = textbox(slide, Inches(0.45), 0, SW - Inches(1.2), Inches(1.0), MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    _set_font(r, size=26, bold=True, color=WHITE)
    if num:
        tf2 = textbox(slide, SW - Inches(1.4), 0, Inches(1.1), Inches(1.0), MSO_ANCHOR.MIDDLE)
        tf2.paragraphs[0].alignment = PP_ALIGN.RIGHT
        r2 = tf2.paragraphs[0].add_run(); r2.text = num
        _set_font(r2, size=14, bold=False, color=RGBColor(0xC8, 0xD4, 0xE6))


def bullets(slide, items, l, t, w, h, size=18, gap=8):
    """items: list of (text, level, color) or str."""
    tf = textbox(slide, l, t, w, h)
    first = True
    for it in items:
        if isinstance(it, str):
            text, lvl, color = it, 0, DARK
        else:
            text, lvl, color = (it + (0, DARK))[:3] if False else (it[0], it[1] if len(it) > 1 else 0, it[2] if len(it) > 2 else DARK)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.level = lvl
        bullet = "• " if lvl == 0 else "– "
        r = p.add_run(); r.text = ("    " * lvl) + bullet + text
        _set_font(r, size=size - 2 * lvl, color=color)
    return tf


def image_fit(slide, name, l, t, w, h):
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    box_ratio = (w / h)
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        nw = w; nh = int(w / img_ratio)
    else:
        nh = h; nw = int(h * img_ratio)
    nl = l + (w - nw) // 2
    nt = t + (h - nh) // 2
    slide.shapes.add_picture(path, nl, nt, width=nw, height=nh)


def caption(slide, text, t):
    tf = textbox(slide, Inches(0.5), t, SW - Inches(1.0), Inches(0.4))
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run(); r.text = text
    _set_font(r, size=12, color=GRAY)


# ---- 幻灯片 1:标题 ----
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, WHITE)
rect(s, 0, Inches(2.7), SW, Inches(0.06), ACC)
tf = textbox(s, Inches(1), Inches(1.5), SW - Inches(2), Inches(1.3), MSO_ANCHOR.MIDDLE)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
r = tf.paragraphs[0].add_run(); r.text = "SO-ARM101 抓取–放置:强化学习 + 视觉"
_set_font(r, size=36, bold=True, color=NAVY)
tf = textbox(s, Inches(1), Inches(3.0), SW - Inches(2), Inches(1.0), MSO_ANCHOR.MIDDLE)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
r = tf.paragraphs[0].add_run(); r.text = "Isaac Lab · PPO(RSL-RL)· 面向真机 SO101 的视觉抓取"
_set_font(r, size=20, color=GRAY)
tf = textbox(s, Inches(1), Inches(4.6), SW - Inches(2), Inches(0.6), MSO_ANCHOR.MIDDLE)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
r = tf.paragraphs[0].add_run(); r.text = "组会汇报   ·   2026-06"
_set_font(r, size=16, color=ACC)


def content_slide(title, num):
    s = prs.slides.add_slide(BLANK)
    title_bar(s, title, num)
    return s


# ---- 2:背景与目标 ----
s = content_slide("一、项目背景与目标", "2 / 14")
bullets(s, [
    ("任务:让 SO-ARM101(5 自由度机械臂)把桌面上的方块抓起,搬运并放到指定目标区", 0),
    ("最终交付:迁移到真实 SO101 机械臂,用摄像头视觉定位方块(不是仿真真值)", 0),
    ("技术栈:Isaac Lab(物理引擎 PhysX 5,GPU)+ RSL-RL(PPO)+ Gymnasium + Python/uv", 0),
    ("★关键认知:训练时用方块真值坐标(无相机),推理时才用视觉", 0, ACC),
    ("→ 策略网络只认坐标、不认图像;视觉只是推理时把图像翻译成坐标的一座桥", 1, GRAY),
], Inches(0.7), Inches(1.4), SW - Inches(1.4), Inches(5.5), size=20, gap=14)

# ---- 3:总体架构 fig1 ----
s = content_slide("二、总体架构:两条流共用同一策略网络", "3 / 14")
image_fit(s, "fig1_overview.png", Inches(0.6), Inches(1.2), SW - Inches(1.2), Inches(5.2))
caption(s, "训练流与推理流唯一的本质区别 = 方块坐标的来源;训练产出 model_26994.pt,推理加载它", Inches(6.6))

# ---- 4:仿真环境与 MDP ----
s = content_slide("三、仿真环境与 MDP 定义", "4 / 14")
bullets(s, [
    ("观测 ≈ 28 维:关节角6 + 关节速6 + 方块xyz真值3 + 目标位姿7 + 上步动作6", 0),
    ("动作 = 6 维:5 个手臂关节目标角(位置控制)+ 夹爪二值开/合", 0),
    ("节拍:物理 100Hz,策略控制 50Hz,每回合 5 秒 = 250 步", 0),
    ("终止:超时(5秒)或 方块掉到桌下", 0),
    ("物理引擎:NVIDIA PhysX 5(GPU),并行跑 4096 个环境——这是高速训练的关键", 0, ACC),
    ("PhysX 与 Isaac Lab 深度绑定,不能换引擎,只能调参(步长/摩擦/求解器)", 1, GRAY),
], Inches(0.7), Inches(1.4), SW - Inches(1.4), Inches(5.5), size=20, gap=13)

# ---- 5:奖励设计 fig6 ----
s = content_slide("四、奖励设计:分阶段门控", "5 / 14")
image_fit(s, "fig6_rewards.png", Inches(0.6), Inches(1.15), SW - Inches(1.2), Inches(5.3))
caption(s, "四阶段(接近抓取→搬运→下降→释放)由高度+距离门控;锁存器防「不抬直接拖到目标」作弊", Inches(6.6))

# ---- 6:PPO 训练 fig7 ----
s = content_slide("五、PPO 训练循环与超参", "6 / 14")
image_fit(s, "fig7_ppo.png", Inches(0.6), Inches(1.15), SW - Inches(1.2), Inches(5.3))
caption(s, "采样→算优势(GAE)→多轮 clip 更新→自适应学习率;每个超参标注在其起作用的位置", Inches(6.6))

# ---- 7:训练日志解读 ----
s = content_slide("六、训练日志怎么看", "7 / 14")
bullets(s, [
    ("Mean reward:每步平均总奖励(所有奖励项加权和)——整体趋势", 0),
    ("Episode_Reward / lifting_object:抓取健康度 ★止损线盯它★(腰斩下跌=死亡螺旋)", 0, RED),
    ("Mean episode length = 250:满步=几乎都正常超时,没有中途掉方块崩盘", 0),
    ("Episode_Reward / 各项:每个奖励项在一回合内的累计贡献", 0),
    ("Curriculum / *:随训练进度变化的惩罚权重;Metrics:目标跟踪误差(非奖励)", 0),
    ("迭代号很大(2.7万)是因为从 model_26994 续训,计数接着加", 1, GRAY),
], Inches(0.7), Inches(1.4), SW - Inches(1.4), Inches(5.5), size=19, gap=13)

# ---- 8:推理工作流 fig3 ----
s = content_slide("七、推理工作流:50Hz 实时闭环", "8 / 14")
image_fit(s, "fig3_inference.png", Inches(0.6), Inches(1.15), SW - Inches(1.2), Inches(5.3))
caption(s, "每步:拍图→识别方块→写入观测→policy(obs)→执行→重拍;实时闭环=掉落可重抓", Inches(6.6))

# ---- 9:视觉感知 颜色掩码 fig5 ----
s = content_slide("八、视觉感知:颜色掩码(采用方案)", "9 / 14")
image_fit(s, "fig5_colormask.png", Inches(0.5), Inches(1.3), SW - Inches(1.0), Inches(4.3))
bullets(s, [
    ("逐像素判红 R>110 且 R>G+35 且 R>B+35 → 取最大连通域 → 求质心 → 针孔反投影", 0),
    ("不训练、纯几何;红方块=画面唯一醒目纯色,色相对光照稳健 → 光照鲁棒性零衰减", 0, GREEN),
], Inches(0.7), Inches(5.7), SW - Inches(1.4), Inches(1.4), size=17, gap=8)

# ---- 10:ResNet 对比 fig4 ----
s = content_slide("九、对比方案:ResNet 监督学习", "10 / 14")
image_fit(s, "fig4_resnet.png", Inches(0.5), Inches(1.3), SW - Inches(1.0), Inches(4.3))
bullets(s, [
    ("监督学习:仿真器自动生成标签(方块真值坐标,免费);数据集上 corr≈0.92(确实在定位)", 0),
    ("但部署存在训练图≠播放图的域差 → 退化为猜均值(~6%) → 被颜色掩码方案取代", 0, ORANGE),
], Inches(0.7), Inches(5.7), SW - Inches(1.4), Inches(1.4), size=17, gap=8)

# ---- 11:实验结果 ----
s = content_slide("十、实验结果(评估漏斗)", "11 / 14")
bullets(s, [
    ("抓取成功率:GT 模式 35.6%  vs  vision 模式 32.0%(感知几乎不丢分)", 0, ACC),
    ("端到端放置(落桌且 XY<8cm):约 8–11%", 0),
    ("策略抓取方式 = 「推-追-夹」:先碰到方块、推动它,再追上去夹住", 0),
    ("评估用自写漏斗(XY距离+是否落桌),而非内置 3D success——因目标点悬空,3D恒为0", 0, GRAY),
    ("结论:视觉感知已不是瓶颈(vision≈GT);瓶颈在抓取率本身 + 放置精度(散布~10cm)", 0),
], Inches(0.7), Inches(1.4), SW - Inches(1.4), Inches(5.5), size=19, gap=14)

# ---- 12:鲁棒性测试 ----
s = content_slide("十一、鲁棒性测试(medium 档)", "12 / 14")
bullets(s, [
    ("光照 ✅ 免疫:29.6%(基线 32%);颜色掩码误差零恶化(红色色相稳)", 0, GREEN),
    ("杂物 ⚠️ 中度:21.4%;静态障碍物遮挡左/中区接近路径", 0, ORANGE),
    ("目标点 ❌ 泛化悬崖:目标出训练框 → 搬运失效(方块被搬回训练目标的均值位置)", 0, RED),
    ("全局摩擦 ❌ 致命:2.8%(μ≈0.35);桌面变滑,「推-追」彻底追不上", 0, RED),
    ("对真机的提示:真实桌面摩擦不可控,是最高风险;放置区应设在训练框内", 0, GRAY),
], Inches(0.7), Inches(1.4), SW - Inches(1.4), Inches(5.5), size=19, gap=14)

# ---- 13:训练侧尝试与结论 ----
s = content_slide("十二、训练侧尝试与结论", "13 / 14")
bullets(s, [
    ("目标:从 v9 微调提升抓取率(治「推-追失败」和「右侧拖动」)", 0),
    ("v17(加推动惩罚)/ v18(加锁存)/ v19(降LR、去噪声)——三次微调均失败", 0, RED),
    ("共同现象:抓取健康度 lifting_object 从 ~0.8 一路阴跌,曲线形状一致", 0),
    ("根因:改过奖励地形后,收敛的 v9 在任何学习率下都不可微调(critic 失配)", 0, ACC),
    ("决策:冻结 v9(model_26994)+ 颜色掩码视觉(抓取 32%)作为当前成果", 0),
], Inches(0.7), Inches(1.4), SW - Inches(1.4), Inches(5.5), size=19, gap=14)

# ---- 14:总结与下一步 ----
s = content_slide("十三、总结与下一步", "14 / 14")
bullets(s, [
    ("已打通完整链路:仿真训练 → 视觉感知 → 推理实时闭环", 0, ACC),
    ("视觉感知已解决:vision 抓取 32% ≈ GT 35.6%,差距不在感知", 0, GREEN),
    ("当前瓶颈:抓取率本身 + 放置精度(~10cm)+ 摩擦鲁棒性", 0),
    ("下一步:换物体 / 换材质测试;真机迁移验证视觉管线", 0),
    ("已沉淀文档:项目流程详解.md(16章)+ 数据流框图.md(7张图)", 0, GRAY),
], Inches(0.7), Inches(1.4), SW - Inches(1.4), Inches(5.5), size=20, gap=15)

out = os.path.join(HERE, "SO-ARM101_组会汇报.pptx")
prs.save(out)
print(f"saved {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
